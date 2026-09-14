"""Build the traceable dissertation-remediation checklist in DOCX and Markdown."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "revision_records"
DOCX_PATH = OUT_DIR / "dissertation_review_remediation_checklist_2026-09-01.docx"
MD_PATH = OUT_DIR / "dissertation_review_remediation_checklist_2026-09-01.md"


SECTIONS = [
    ("1. 核心 provenance、泄漏与复现性问题", [
        ("A01", "已完成", "Polyvore 实际来源与许可边界", "Methodology 3.2、Appendix D.1、references.bib：补作者公开 repository、获取日期和本地 Maryland metadata repository commit；明确代码许可不等于第三方图片再分发许可。"),
        ("A02", "已完成", "Zara 来源不完整", "Methodology 3.2、Appendix D.1：记录 Kaggle 的 Zara Products、作者、URL、获取日期与页面所示 MIT 标签；仅用于识别准备和 feature-domain diagnostic，不把该标签扩大解释为品牌图片权利。"),
        ("A03", "已完成", "SigLIP checkpoint 不可复现", "Methodology 3.3、Appendix D.1：写明 google/siglip-base-patch16-224、本地 cache revision 7fd15f0…2761ed、Transformers 版本。"),
        ("A04", "已完成", "processed counts 形成过程", "Methodology 3.2、Appendix D.2：列出 source rows、retained、positive、negative、insufficient supported positions 与 duplicate-position audit；说明 duplicate 是已排除行的重叠原因，不重复扣除。"),
        ("A05", "已完成", "negative sampling 是否 split-local", "Methodology 3.2：明确每个 split 从各自 published compatibility file 独立准备；未跨 split 构造替换项。"),
        ("A06", "已纠正", "disjoint 属性未经核验", "新增 recommendation_training/audit_split_overlap.py、tests/test_split_overlap_audit.py 与 results/polyvore_disjoint_overlap_audit.json。结果：outfit ID overlap 为 0；item ID overlap 为 train–valid 3,781、train–test 84、valid–test 34。全文不再宣称严格 item-disjoint。"),
        ("A07", "已完成", "FITB identifier matching", "Methodology 3.4、tests/test_official_fitb.py：说明由 set/item identifier 恢复 shuffled candidate 的正确答案；测试覆盖 shuffled order 与 slot mismatch。"),
        ("A08", "已完成", "1,670 abstract states 构造不清", "Methodology 3.5、Appendix D.3：仅使用 training division，按 slot×type×colour×style×audience 分组，取离 centroid 最近的 observed embedding 作为 medoid；外套可由天气约束 mask。"),
        ("A09", "已完成", "379,265 参数缺少可计算规格", "Methodology 3.3、Appendix D.3：补 768→256 projection、4×32 position embeddings、288-dimensional global/pairwise summaries、576→256→128→1 scorer、dropout 0.20 及逐项参数。"),
        ("A10", "已完成", "bootstrap 协议不完整", "Methodology 3.4、Appendix D.3：记录 95% percentile interval、resampling unit 和次数；AUC 1,000 rows、FITB 2,000 questions、paired ablation 5,000 questions、search score loss 2,000 queries。"),
        ("A11", "已完成", "ECE 设置缺失", "Methodology 3.4、Appendix D.3：validation 上 10 个 equal-width bins；temperature 为 validation NLL optimisation 的单一标量。"),
        ("A12", "已完成", "exact-search safety-limit 缺少证据", "Appendix D.3、tests/test_abstract_search.py：超出安全上限时显式抛出 RuntimeError，不把 prefix 结果称为 exact。"),
        ("A13", "已完成", "ethics waiver 措辞与证据", "Appendix A 已嵌入三项真实培训证明；Appendix B 已嵌入 2026 年 5 月 12 日的真实 waiver confirmation。正文采用文件中的 waived / does not require approval 表述，并说明证明使用的是项目早期简短标题。"),
        ("A14", "已完成", "mean-pooled baseline 未在 Results 报告", "Results 4.2 已加入 test AUC 0.5532；Discussion 中仅作为 controlled baseline 使用。"),
        ("A15", "已完成", "0.8573 与 0.8583 易混", "Results 4.2 明确 0.8573 是 fixed-seed controlled baseline run，0.8583 是 selected checkpoint final test result。"),
        ("A16", "已完成", "0.6025 与 0.5994 易混", "Results 4.4 明确 0.6025 是 selected seed-42 checkpoint，0.5994 是 three-seed full-model mean。"),
    ]),
    ("2. 引用、来源与参考文献", [
        ("B01", "已完成", "补 Zara data source", "references.bib 增加 Kaggle 数据条目；正文及 Appendix 通过 author–date citation 链接到原页面。"),
        ("B02", "已完成", "补 Polyvore artefact source", "references.bib 增加作者公开 dataset/code repository；与 Vasileva et al. 学术论文条目并列，分别承担方法和 artefact provenance。"),
        ("B03", "已完成", "补 Open-Meteo", "references.bib 与 Appendix D.1 增加官方 API documentation 和访问日期。"),
        ("B04", "已完成", "软件栈版本", "Appendix D.4 记录 Python、NumPy、pandas、Pillow、scikit-learn、PyTorch、Transformers 与 Streamlit 的实际环境版本。"),
        ("B05", "已完成", "SigLIP model-card provenance", "references.bib 增加 Google model card；Zhai et al. 仍用于方法背景，model card 用于具体 checkpoint。"),
        ("B06", "已完成", "Bouthillier author metadata", "references.bib 将截断作者修正为 Nazanin Mohammadi Sepahvand。"),
        ("B07", "已完成", "ACM article number", "Deldjoo、Mehrabi、Keyes 分别补 Article 87、115、88 与相应页码/文章长度信息。"),
        ("B08", "不采纳原建议", "统一为 Manchester Harvard", "项目属于 Warwick，不应无依据切换 Manchester 规范。main.tex 保持一致 author–date 样式并修正元数据、DOI/URL；最终仍应以 WMG 当前 handbook/template 的引用要求为准。"),
        ("B09", "已完成", "DOI/URL 断行与可访问性", "hyperref 保留可点击链接；长 commit/revision 从窄表格移到正文并使用可断行显示，避免表格溢出。"),
        ("B10", "已完成", "Selbst/Mehrabi citation fit", "偏差来源由 Mehrabi 支撑；Selbst 只用于更广泛的社会技术限制，不用它直接证明 frozen representation bias。"),
        ("B11", "已完成", "Zhai claim scope", "正文仅用 SigLIP 文献解释预训练表示和方法来源，不把它写成 fashion-domain transfer 已被证明。"),
        ("B12", "已完成", "same-category negatives 比较性断言", "Discussion 改为 ‘reduces trivial category-based distinctions’，不再声称未经直接对照实验支持的普遍难度提升。"),
    ]),
    ("3. Ethics、Responsible AI 与权利边界", [
        ("C01", "已完成", "伦理审批术语", "Appendix B 的原始确认明确写明 ethical approval be waived，并说明项目 does NOT require approval；Introduction 与附录说明已据此统一，未虚构 reference number。"),
        ("C02", "已核对", "gender inference", "系统保留用户手动 menswear/womenswear candidate-range 选择，不从图像自动推断人口属性；论文不把该设计写成 fairness 已被证明。"),
        ("C03", "已核对", "bias 表述", "Discussion 6.7/6.8 将风险限定为 data selection、labels、pretrained representation、candidate construction 与 interaction design；无 demographic fairness 结论。"),
        ("C04", "已核对", "用户照片与本地数据", "Methodology 3.8 与本地数据图：用户照片、wardrobe JSON、saved cities 为 local-first state，不进入研究训练集；天气服务仅接收地点相关请求。"),
        ("C05", "已保留", "future participant study", "Discussion future work 保留 consent、withdrawal、secure handling 和 ethics approval，明确是未来研究要求。"),
        ("C06", "已压缩", "商业化扩展", "删除 conversion/retention/willingness-to-pay 等无证据讨论，仅保留一句 prospective retailer/independent-designer linking，并明确需权限、实时数据与独立商业评估。"),
        ("C07", "需要本人确认", "生成式 AI 使用声明", "未擅自写入最终 submission。学生必须按 WMG assessment brief 向导师/module team 确认允许范围，并如实说明 why/where/how；现有审计草案位于 word_count_and_ethics_audit_2026-09-01.md。"),
    ]),
    ("4. 结果—方法证据链", [
        ("D01", "已完成", "Calibration", "Results 报告 Brier/NLL/ECE/temperature scaling；Methodology/Appendix 补 ECE 实现。"),
        ("D02", "已完成", "MRR/NDCG", "Results 4.4 保留 MRR 0.7705、NDCG 0.8289，不再当作缺失项。"),
        ("D03", "已完成", "FITB coverage", "Results 报告 4,468/15,145（29.50%）及过滤原因，结论仅适用于 retained subset。"),
        ("D04", "已完成", "menswear sparsity", "Appendix C.1 保留 46 vs 4,422 的 descriptive coverage；不作稳定 subgroup/fairness 推断。"),
        ("D05", "已完成", "pairwise ablation", "Analysis 5.3 保留 −10.28 percentage points 及 paired CI；作为最强结构性证据。"),
        ("D06", "已完成", "position embedding", "Analysis 5.3 明确 CI crossed zero；保留该组件只是低成本的 explicit slot cue，不称为 empirically necessary。"),
        ("D07", "已完成", "exact vs quota", "Results/Analysis 报告 top-one recovery、top-ten recall、runtime 与 score loss；18-query 选择在 Methodology 说明为 deterministic engineering coverage。"),
        ("D08", "已完成", "feature-domain diagnostic", "AUC 1.000 只解释为 sampled feature distributions readily distinguishable；不推导 Zara compatibility performance drop。"),
        ("D09", "已完成", "overfitting", "Analysis 5.4 同时报告 train AUC 0.9699 与 test 0.8583 的 gap，以及 validation/test 接近；不使用 ‘no overfitting’。"),
        ("D10", "无法完全追溯", "position retention 是否 test-informed", "现有 artefact 时间与文本不足以证明最初决定发生在 test inspection 前。论文已避免把 retention 写成 test-validated selection，并把无 position embedding 作为合理未来替代；若导师要求严格 chronology，学生需提供实验日志。"),
        ("D11", "已修正", "ranking-loss causal overreach", "Discussion 6.4 改成 under the tested setup did not improve the retained evaluation，不再把原因归因于未经隔离验证的 candidate construction/representation limits。"),
        ("D12", "已修正", "attention model size", "Analysis 5.3 删除 ‘much larger’，改为不需要 attention-based interaction module。"),
    ]),
    ("5. 主动防御与重复性清理", [
        ("E01", "已删除", "UI emotional-state denial", "Methodology 3.7 只保留 visual personalisation；删除 ‘not evidence that interface changes emotional state’。"),
        ("E02", "已压缩", "privacy/provenance 并列防御", "Methodology 3.8 分开说明 local-first privacy 与 data provenance，不再用无必要的 but 句式强行绑定。"),
        ("E03", "已删除", "fallback stylistic-correctness 重复", "Methodology 3.8、Analysis 5.6、Discussion 6.5 中删除重复否认，只保留 learned/fallback evidence 不可混用的技术边界。"),
        ("E04", "已压缩", "user liking/fashion/purchase 重复", "Analysis 5.1/5.2 与 Discussion 6.1/6.3 删除或合并；participant limitation 集中在 Analysis 5.8 与 Discussion validity/future work。"),
        ("E05", "已压缩", "universal superiority 重复", "Analysis 5.2、5.3 与 Discussion 6.2 仅保留 controlled adaptations 不能等同 original-system reproduction 的必要边界。"),
        ("E06", "已删除", "SigLIP pretraining cost 重复", "轻量化定义保留在 Introduction/Methodology，Chapter 5 和 Conclusion 不再重复同一句。"),
        ("E07", "已删除", "possible explanation disclaimer", "Analysis 5.3 用自然限定语 ‘One possible explanation’ 表达，不再追加机械性否认句。"),
        ("E08", "已压缩", "general reliability 防御", "Analysis 5.4、Conclusion 改为直接陈述 held-out validation/test consistency 与 train–test gap。"),
        ("E09", "已删除", "defensible stopping point", "删除主观性收尾；仅报告 validation-based selection 与 test evidence。"),
        ("E10", "已压缩", "weighting 永久无效防御", "Analysis/Discussion 只说 tested weighting schemes 未达到预设 acceptance gate。"),
        ("E11", "已保留并正向改写", "procedural explanation boundary", "Analysis 5.6 保留 system-level procedural explanation，因为这是解释性声明的必要边界；改为正向定义其覆盖范围。"),
        ("E12", "已保留", "AUC 不覆盖规则失败", "Analysis 5.6 保留：这是区分 learned ranking metric 与 deterministic constraint tests 的必要 evaluation boundary。"),
        ("E13", "已保留", "search scalability boundary", "Analysis 5.7 保留 recorded hardware 与 bounded search space，因为 runtime 结论必须限定测量条件。"),
        ("E14", "已保留", "domain-shift interpretation", "Analysis 5.8 保留：domain classifier AUC 不能换算成 compatibility accuracy drop，属于必要统计解释。"),
        ("E15", "已删除", "shops/cultural appropriateness 新防御", "Discussion 6.5 删除未研究且会扩展范围的 shops、universal cultural appropriateness 和重复 stylistic correctness 句。"),
        ("E16", "已删除", "Chapter 6 Summary 重复", "删除原 6.11 Summary；Chapter 7 承担最终回答、贡献与限制。"),
    ]),
    ("6. Chapter 6 结构决定", [
        ("F01", "已完成", "6.1 Introduction", "缩短为 RQ 与 claim-level 导向，删除 prefer/trust/purchase 串联防御。"),
        ("F02", "已完成", "6.2 previous research", "压缩 controlled-adaptation 与 universal-superiority 重复，仅保留与文献的实质比较。"),
        ("F03", "已完成", "6.3–6.5 RQ answers", "保留核心 evidence；删除 AUC=percentage-correct、human approval、shops 与 fallback-style 等重复说明。"),
        ("F04", "保留但压缩", "6.6 objectives mapping", "未按清单整节删除。理由：导师明确要求 Discussion 与 objectives link，Warwick learning outcomes 也要求说明目标完成情况；现压成单段，不再逐项复述所有数字。"),
        ("F05", "保留并收紧", "6.7 validity", "未删除。理由：导师明确肯定 internal/external/construct/ecological validity，且项目要求评价 findings confidence。已加入 item-overlap 审计并减少重复。"),
        ("F06", "已完成", "6.8 contribution/commercial", "保留研究贡献与 responsible-use 内容；商业化只留一段短的 prospective route。"),
        ("F07", "已完成", "6.9 applied-AI implications", "保留新增 synthesis：task-matched architecture、data transformation、retrieval bottleneck、auditability。"),
        ("F08", "已完成", "6.10 future work", "按 evidence gaps 排序：licensed data、target-domain validation、ethics-approved participant study、search scaling。"),
        ("F09", "已删除", "6.11 summary", "删除与 Analysis summary 和 Conclusion 重复的段落。"),
    ]),
    ("7. Appendix、图表与实现一致性", [
        ("G01", "已完成", "Appendix D provenance/reproducibility", "新增 provenance、processing audit、split overlap、model parameter、bootstrap/ECE、FITB/search tests、candidate construction 与 software versions。"),
        ("G02", "已核对", "Recommendation UML", "CompatibilityRanker、OutfitCandidateRanker、ModelConfig、LightweightOutfitRanker 以及 function-based adapter/facade 均能对应当前代码。图中已把 adapter/facade 标为模块职责而非伪造 class。"),
        ("G03", "已核对", "Persistence ERD", "UserPreferences 为独立 local file；saved cities 上限为 6；WardrobeStore 默认组为 Everyday/Work/Occasion/Travel，因此 4..* defaults plus custom groups 与代码一致。"),
        ("G04", "已核对", "Acceptance evidence chain", "recognition/FITB/search safety/fallback/local persistence/delete confirmation 均有实现或测试路径；Appendix D 指向具体 artefact/test，不把 UML 本身当成测试证据。"),
        ("G05", "已修正", "模型架构图重叠", "model_architecture.tex 已重排节点和箭头，避免原 Figure 3.2 文本与连线重叠；最终 PDF 需随每次正文更新重新视觉复核。"),
    ]),
]


OPEN_ITEMS = [
    "student name、student ID 与 submission date 已填写；仍须由学生提供准确 qualification 并完成 signature。",
    "从 MSc Applied Artificial Intelligence Specific Course Regulations 粘贴准确 project definition，并完成 relationship statement；不得根据课程网页自行拼写定义。",
    "向导师或 WMG module team 确认本 assessment 的生成式 AI 允许范围和声明格式；如实披露实际使用。",
    "若学校要求证明 position-embedding 决策未受 test set 影响，需提供当时的 dated experiment log；现有材料不足以复原决策时间线。",
    "第三方图片可公开访问不等于可再分发；公开 repository 前继续排除 Polyvore/Zara 源图、item-level derivatives，并让导师/学校确认 checkpoint 发布边界。",
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row_together(row, repeat_header=False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if repeat_header:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_run_font(run, name="Calibri", size=Pt(10.5), bold=False, colour=None) -> None:
    run.font.name = name
    run.font.size = size
    run.bold = bold
    if colour:
        run.font.color.rgb = RGBColor(*colour)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, end))
    set_run_font(run, size=Pt(9), colour=(90, 102, 116))


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_page_number(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, colour, before, after in (
        ("Title", 19, (31, 77, 120), 0, 14),
        ("Heading 1", 16, (46, 116, 181), 18, 10),
        ("Heading 2", 13, (46, 116, 181), 14, 7),
        ("Heading 3", 12, (31, 77, 120), 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*colour)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_status_table(doc: Document, rows) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    table.style = "Table Grid"
    widths = (Inches(0.55), Inches(0.85), Inches(1.65), Inches(3.45))
    headers = ("ID", "状态", "审查问题", "整改、位置与证据")
    keep_row_together(table.rows[0], repeat_header=True)
    for idx, (cell, label, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        set_run_font(run, size=Pt(9.5), bold=True, colour=(31, 77, 120))
    status_fill = {
        "已完成": "EAF4EA", "已纠正": "EAF4EA", "已修正": "EAF4EA",
        "已核对": "EAF4EA", "已删除": "EAF4EA", "已压缩": "EAF4EA",
        "已保留": "FFF4D6", "已保留并正向改写": "FFF4D6",
        "保留但压缩": "FFF4D6", "保留并收紧": "FFF4D6",
        "部分完成": "FDE9D9", "需要本人材料": "FDE9D9", "需要本人确认": "FDE9D9",
        "无法完全追溯": "FDE9D9", "不采纳原建议": "E8EEF5",
    }
    for item_id, status, issue, evidence in rows:
        row = table.add_row()
        keep_row_together(row)
        cells = row.cells
        values = (item_id, status, issue, evidence)
        for idx, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=Pt(9.1), bold=(idx in (0, 1)))
        set_cell_shading(cells[1], status_fill.get(status, "FFFFFF"))
    doc.add_paragraph()


def build_docx() -> None:
    doc = Document()
    configure_document(doc)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("硕士论文整改与提交前核查清单")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Cove Complete-Outfit Recommendation Dissertation · 2026-09-01")
    set_run_font(run, size=Pt(11), colour=(90, 102, 116))

    box = doc.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    cell = box.cell(0, 0)
    set_cell_shading(cell, "E8F1FA")
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    p = cell.paragraphs[0]
    r = p.add_run("审查结论：核心方法与证据链问题已整改。最重要的新发现是 published ‘disjoint’ split 在本地 metadata 中并非严格 item-disjoint；论文已如实披露剩余 item overlap 和 memorisation risk。真实伦理培训证明与 waiver confirmation 已补齐。稿件仍不能称为可直接提交，因为 qualification、签名、课程定义/关系声明与生成式 AI 合规确认尚未完成。")
    set_run_font(r, size=Pt(10.5), bold=True, colour=(31, 77, 120))

    doc.add_heading("使用与判定原则", level=1)
    for text in (
        "本清单逐项判断原审查意见是否适用于 Warwick WMG 项目，不机械执行来源不明或与学校不匹配的格式建议。",
        "‘已完成’表示正文/附录/代码或测试已有可定位证据；‘部分完成/需要本人材料’表示不能在没有真实文件或本人信息的情况下代填。",
        "主动防御语句已删除或合并；只有会改变指标含义、研究范围、许可边界或证据强度的限定语被保留。",
        "文件位置以 section 与稳定路径为主，避免 LaTeX 修改后行号漂移。",
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)

    doc.add_heading("总体状态", level=1)
    summary = doc.add_table(rows=4, cols=2)
    summary.style = "Table Grid"
    for row, (label, value) in enumerate((
        ("正文估算字数", "15,567（Abstract + Chapter 1–7；目标 15,000，处于 ±10% 范围）"),
        ("自动化检查", "26/26 citation keys 可解析；11 项新增/关键单元测试通过；PDF 82 页"),
        ("最终证据状态", "方法、结果、限制与 provenance 基本闭合；严格 item-disjoint claim 已撤回"),
        ("提交阻断项", "qualification、签名、课程定义/关系声明、AI 使用规则确认"),
    )):
        for cell in summary.rows[row].cells:
            set_cell_margins(cell)
        set_cell_shading(summary.cell(row, 0), "E8EEF5")
        summary.cell(row, 0).text = label
        summary.cell(row, 1).text = value
        for idx, cell in enumerate(summary.rows[row].cells):
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=Pt(10), bold=(idx == 0))

    for heading, rows in SECTIONS:
        doc.add_heading(heading, level=1)
        add_status_table(doc, rows)

    doc.add_heading("8. 尚未完成：学生提交前动作", level=1)
    for idx, text in enumerate(OPEN_ITEMS, 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(text)

    doc.add_heading("9. 已执行的自动检查", level=1)
    checks = [
        "Manuscript source check：24 个 TeX 文件、26 bibliography entries、26 citation keys，无缺失引用键。",
        "Targeted unit tests：split-overlap audit、official FITB matching、abstract medoid/exact search/safety gate，共 11 项通过。",
        "PDF integrity：82 页；无 embedded files、invisible-text rendering、optional-content markers 或所列隐藏提示词命中。",
        "PDF word estimate：15,567；脚本 dissertation/tools/count_pdf_words.py 可重复运行。",
        "伦理附件版面已逐页复核；仍需人工完成 qualification、签名、课程定义/关系声明、AI 合规确认及提交前最终通读。",
    ]
    for text in checks:
        doc.add_paragraph(text, style="List Bullet")

    doc.core_properties.title = "Cove Dissertation Remediation Checklist"
    doc.core_properties.subject = "Traceable review and completion record"
    doc.core_properties.author = "Cove project review record"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def build_markdown() -> None:
    lines = [
        "# 硕士论文整改与提交前核查清单",
        "",
        "Cove Complete-Outfit Recommendation Dissertation · 2026-09-01",
        "",
        "> 核心方法与证据链问题已整改。published ‘disjoint’ split 在本地 metadata 中并非严格 item-disjoint；论文已如实披露。真实伦理附件已补齐；qualification、签名、课程定义/关系声明与 AI 合规确认仍需学生完成。",
        "",
    ]
    for heading, rows in SECTIONS:
        lines.extend((f"## {heading}", "", "| ID | 状态 | 审查问题 | 整改、位置与证据 |", "|---|---|---|---|"))
        for item_id, status, issue, evidence in rows:
            lines.append(f"| {item_id} | {status} | {issue} | {evidence} |")
        lines.append("")
    lines.extend(("## 8. 尚未完成：学生提交前动作", ""))
    lines.extend(f"{idx}. {text}" for idx, text in enumerate(OPEN_ITEMS, 1))
    lines.append("")
    MD_PATH.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_markdown()
    print(DOCX_PATH)
    print(MD_PATH)
